"""盤點資料庫裡每張表的欄位結構，存成 Excel。

連線讀 config.json 與 .env，跟 app.pyw、import_data.py 連的是同一個資料庫。

用法：
    python column_inventory.py                       # 全部 schema，存成 欄位結構_資料庫_日期.xlsx
    python column_inventory.py 輸出.xlsx             # 指定檔名
    python column_inventory.py --schemas dbo,prc     # 只盤點這幾個 schema
"""
import sys
from datetime import date

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text

import core

# 型別照 SSMS 的寫法組回來：NVARCHAR 的 max_length 是 byte 數，要除以 2。
QUERY = """
SELECT s.name AS [Schema], t.name AS [表], c.name AS [欄位],
       UPPER(ty.name) +
       CASE
           WHEN ty.name IN ('nvarchar', 'nchar')
               THEN '(' + IIF(c.max_length = -1, 'MAX', CAST(c.max_length / 2 AS varchar(10))) + ')'
           WHEN ty.name IN ('varchar', 'char', 'varbinary', 'binary')
               THEN '(' + IIF(c.max_length = -1, 'MAX', CAST(c.max_length AS varchar(10))) + ')'
           WHEN ty.name IN ('decimal', 'numeric')
               THEN '(' + CAST(c.precision AS varchar(5)) + ',' + CAST(c.scale AS varchar(5)) + ')'
           WHEN ty.name IN ('datetime2', 'time', 'datetimeoffset')
               THEN '(' + CAST(c.scale AS varchar(5)) + ')'
           ELSE ''
       END AS [型別],
       IIF(c.is_nullable = 1, 'Y', 'N') AS [可為NULL]
FROM sys.tables t
JOIN sys.schemas s ON s.schema_id = t.schema_id
JOIN sys.columns c ON c.object_id = t.object_id
JOIN sys.types ty  ON ty.user_type_id = c.user_type_id
WHERE t.is_ms_shipped = 0
ORDER BY s.name, t.name, c.column_id
"""

FONT = "微軟正黑體"           # 跟規格表一致
WIDTHS = {"Schema": 10, "表": 30, "欄位": 24, "型別": 18, "可為NULL": 10}


def inventory(engine, schemas=None):
    """每張表每個欄位一列。schemas 給了就只留那幾個（不分大小寫）。"""
    with engine.connect() as conn:
        df = pd.read_sql(text(QUERY), conn)
    if schemas:
        keep = {s.strip().lower() for s in schemas}
        df = df[df["Schema"].str.lower().isin(keep)]
    return df.reset_index(drop=True)


def save(df, path):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="欄位結構", index=False)
        ws = xw.sheets["欄位結構"]
        head = PatternFill("solid", fgColor="FF1F4E78")
        for cell in ws[1]:
            cell.font = Font(name=FONT, bold=True, color="FFFFFFFF")
            cell.fill = head
            cell.alignment = Alignment(horizontal="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name=FONT)
        for i, name in enumerate(df.columns, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = WIDTHS.get(name, 14)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions


def main(*argv):
    args = list(argv)
    schemas = None
    if "--schemas" in args:
        i = args.index("--schemas")
        if i + 1 >= len(args):
            sys.exit("--schemas 後面要接 schema 名稱，例如 --schemas dbo,prc")
        schemas = args[i + 1].split(",")
        del args[i:i + 2]

    cfg = core.load_config()
    if cfg["conn"].startswith("mssql"):
        missing = core.missing_driver()
        if missing:
            sys.exit(missing)
    engine = core.connect(cfg["conn"])
    try:
        df = inventory(engine, schemas)
        db = engine.url.database or "db"
    except Exception as e:
        sys.exit("連線失敗：" + core.explain_conn_error(e))

    if df.empty:
        sys.exit("沒有找到任何表" + (f"（schema：{','.join(schemas)}）" if schemas else ""))
    path = args[0] if args else f"欄位結構_{db}_{date.today():%Y%m%d}.xlsx"
    save(df, path)
    tables = df.groupby(["Schema", "表"]).ngroups
    print(f"{tables} 張表、{len(df):,} 個欄位 → {path}")
    for schema, n in df.groupby("Schema")["表"].nunique().items():
        print(f"  {schema}: {n} 張")
    return 0


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:]))
