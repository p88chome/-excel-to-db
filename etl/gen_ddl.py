"""從規格表（資料移轉需求規格及欄位屬性.xlsx）產生建表的 T-SQL。

規格改了就重跑一次，DDL 不用手改：
    python etl/gen_ddl.py 規格.xlsx > etl/sql/01_create_tables.sql

只讀「(SAP) schema.Table」這種分頁。灰底刪除線的列（規格上標為移除）不建。
規格上的 VARCHAR 一律建成 NVARCHAR，長度不變（NVARCHAR(10) 一樣是 10 個字）。
"""
import re
import sys

import openpyxl

SHEET_RE = re.compile(r"^\(SAP\)\s*(\w+)\.(\w+)$")


def columns(ws):
    """(欄名, 型別, 是否為鍵)，跳過空白列與刪除線的列。"""
    for r in range(2, ws.max_row + 1):
        field = ws.cell(r, 2).value
        if not field or ws.cell(r, 2).font.strike:
            continue
        spec = str(ws.cell(r, 5).value).strip().upper()
        # 規格寫 VARCHAR 的一律建成 NVARCHAR：stg 表全是 NVARCHAR，型別一致
        # 才不會在 JOIN 時隱含轉換；中文、日文也不用一欄一欄判斷。
        spec = re.sub(r"^VARCHAR\b", "NVARCHAR", spec)
        yield str(field).strip(), spec, ws.cell(r, 7).value == "Y"


def ddl(schema, table, cols):
    keys = [c for c, _, k in cols if k]
    width = max(len(c) for c, _, _ in cols) + 2
    lines = [f"    {('[' + c + ']').ljust(width)} {t}{' NOT NULL' if k else ' NULL'}"
             for c, t, k in cols]
    if keys:
        lines.append(f"    CONSTRAINT PK_{schema}_{table} PRIMARY KEY ("
                     + ", ".join(f"[{k}]" for k in keys) + ")")
    return (f"IF OBJECT_ID(N'{schema}.{table}', N'U') IS NOT NULL DROP TABLE {schema}.{table};\n"
            f"CREATE TABLE {schema}.{table} (\n" + ",\n".join(lines) + "\n);\nGO\n")


def main(path):
    wb = openpyxl.load_workbook(path)
    tables = []
    for ws in wb.worksheets:
        m = SHEET_RE.match(ws.title)
        if m:
            tables.append((m.group(1).lower(), m.group(2), list(columns(ws))))

    out = [f"-- 由 etl/gen_ddl.py 從規格表產生，不要手改。\n-- 來源：{path.split(chr(92))[-1]}\n",
           "-- 注意：會 DROP 掉同名的表重建，表裡的資料會不見。\n"]
    for schema in sorted({s for s, _, _ in tables}):
        out.append(f"IF SCHEMA_ID(N'{schema}') IS NULL EXEC(N'CREATE SCHEMA {schema}');\nGO\n")
    out.append("IF SCHEMA_ID(N'etl') IS NULL EXEC(N'CREATE SCHEMA etl');\nGO\n")
    for schema, table, cols in tables:
        out.append(ddl(schema, table, cols))
    sys.stdout.reconfigure(encoding="utf-8-sig")
    print("\n".join(out))


if __name__ == "__main__":
    main(sys.argv[1])
